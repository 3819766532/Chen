# ============================================================
# 免责声明：本代码仅供学习交流使用，请勿用于非法用途。
# 使用本代码所产生的一切后果由使用者自行承担，
# 作者不承担任何法律责任。
# ============================================================

import openpyxl
from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.title = "测试数据"
ws.append(["姓名", "年龄", "城市"])
data = [
    ["张三", 25, "北京"],
    ["李四", 30, "上海"],
    ["王五", 28, "广州"]
]
for row in data:
    ws.append(row)
print(f"读取到的第一个数据是: {ws['A1'].value}")
print(f"读取到的第二行数据是: {ws['A2'].value}, {ws['B2'].value}, {ws['C2'].value}")
print("\n--- openpyxl 测试通过！Excel 功能已就绪 ---")
